"""
WSR Admin Data Import
=====================
Loads the source-of-truth Excel workbooks into the PRIVATE admin tables
(ds_*) used by the hidden admin dashboard. Run scripts/admin_schema.sql in the
Supabase SQL Editor first.

Reads directly from disk (nothing is copied into this public repo):
    • DATA STORAGE & ANALYSIS 2026.xlsx  → ds_categories, ds_entries, ds_final_rank
    • FINALRANKING2026.xlsx              → ds_settings (latest sports added),
                                           plus a cross-check of the 2026 totals
    • COUNTRIESLIST.xlsx                 → ds_countries (canonical 206 NOCs)

Requirements (once):
    .venv/bin/pip install supabase openpyxl

Usage:
    # Parse + validate only, no upload:
    .venv/bin/python scripts/import_admin_data.py --dry-run

    # Full load (REPLACES the ds_* tables with what's in the workbooks):
    SUPABASE_SERVICE_KEY="sb_secret_..." .venv/bin/python scripts/import_admin_data.py

The workbook folder defaults to ../Ranking 26 (next to this repo); override
with WSR_DATA_DIR=/path/to/folder. The year in the file names (2026) can be
changed with WSR_DATA_YEAR=2027.

⚠ Re-importing overwrites edits made in the dashboard. If the dashboard has
  recorded changes since the last import the script stops and asks for --force.
  Export from the dashboard (Entries → Export CSV) first if you want to keep them.
"""

import collections
import datetime as dt
import json
import os
import sys
from pathlib import Path

# ── Configuration ─────────────────────────────────────────────────────────────
SUPABASE_URL = os.environ.get("SUPABASE_URL", "https://sliiydlqtuopadtbrwtx.supabase.co")
# Must be the service_role SECRET key (the ds_* tables have no public access).
# Never hardcode or commit it — pass it at runtime.
SUPABASE_KEY = os.environ.get("SUPABASE_SERVICE_KEY")

DATA_YEAR = os.environ.get("WSR_DATA_YEAR", "2026")
DATA_DIR = Path(os.environ.get("WSR_DATA_DIR", Path(__file__).parent.parent.parent / "Ranking 26"))
STORAGE_FILE = DATA_DIR / f"DATA STORAGE & ANALYSIS {DATA_YEAR}.xlsx"
FINAL_FILE = DATA_DIR / f"FINALRANKING{DATA_YEAR}.xlsx"
COUNTRIES_FILE = DATA_DIR / "COUNTRIESLIST.xlsx"
# ─────────────────────────────────────────────────────────────────────────────

BATCH_SIZE = 1000

# Spellings found in the workbooks → canonical COUNTRIESLIST name, shared with
# the dashboard's "Upload Excel" tab (lib/admin/countryAliases.json). Excel's
# lookups are case-insensitive so mixed case already matches; these are the
# typos / old names that don't. Names not listed there and not in COUNTRIESLIST
# (e.g. CURAÇAO, BONAIRE — not NOCs) are kept as-is with no country code.
ALIASES = {
    k: v for k, v in json.loads(
        (Path(__file__).parent.parent / "lib" / "admin" / "countryAliases.json").read_text(encoding="utf-8")
    ).items() if not k.startswith("_")
}


def load_workbook(path: Path):
    try:
        import openpyxl
    except ImportError:
        print("ERROR: openpyxl not installed. Run:  .venv/bin/pip install openpyxl")
        sys.exit(1)
    if not path.exists():
        print(f"ERROR: workbook not found: {path}")
        sys.exit(1)
    # data_only=True → cached formula results (what Excel shows), not formulas
    return openpyxl.load_workbook(path, read_only=True, data_only=True)


def to_int(val):
    """Ranks are occasionally stored as text ("4") in the workbook."""
    if val is None:
        return None
    try:
        return int(float(str(val).strip()))
    except ValueError:
        return None


def to_num(val):
    if val is None:
        return None
    try:
        return float(val)
    except (TypeError, ValueError):
        return None


# ── Readers ───────────────────────────────────────────────────────────────────
def read_countries():
    ws = load_workbook(COUNTRIES_FILE).active
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[0] or not r[1]:
            continue
        rows.append({
            "code":           str(r[0]).strip().upper(),
            "name":           str(r[1]).strip().upper(),
            "iso_3":          (str(r[2]).strip().upper() if r[2] else None),
            "iso_2":          (str(r[3]).strip().lower() if r[3] else None),
            "continent_code": (str(r[4]).strip().upper() if r[4] else None),
        })
    return rows


class CountryResolver:
    def __init__(self, countries):
        self.by_name = {c["name"]: c["code"] for c in countries}
        self.unmatched = collections.Counter()

    def __call__(self, raw):
        name = " ".join(str(raw).split()).upper()
        name = ALIASES.get(name, name)
        code = self.by_name.get(name)
        if code is None:
            self.unmatched[name] += 1
        return name, code


def read_storage(resolve):
    wb = load_workbook(STORAGE_FILE)

    # CATEGORIES: columns F (sport), G (type), H (optional note)
    categories = []
    for r in wb["CATEGORIES"].iter_rows(values_only=True):
        if len(r) < 7 or not r[5] or not r[6]:
            continue
        categories.append({
            "sport": " ".join(str(r[5]).split()).upper(),
            "type":  str(r[6]).strip().upper(),
            "note":  (str(r[7]).strip() if len(r) > 7 and r[7] else None),
        })

    # BY COUNTRY: columns A–E (Year, SPORT, RANK, COUNTRY, POINTS). Column F is
    # a VLOOKUP into CATEGORIES — derived, so it isn't stored.
    entries, skipped = [], []
    for i, r in enumerate(wb["BY COUNTRY"].iter_rows(min_row=2, max_col=5, values_only=True), start=2):
        if all(v is None for v in r):
            continue
        year, sport, rank, country, points = r
        year, rank, points = to_int(year), to_int(rank), to_num(points)
        if year is None or not sport or rank is None or not country or points is None:
            skipped.append((i, r))
            continue
        name, code = resolve(country)
        entries.append({
            "year": year, "sport": " ".join(str(sport).split()).upper(),
            "rank": rank, "country": name, "country_code": code, "points": points,
            "updated_by": "import",
        })

    # FINAL RANK: columns A–E (year, rank, country, points, progress)
    final_rank, fr_skipped = [], []
    for i, r in enumerate(wb["FINAL RANK"].iter_rows(min_row=2, max_col=5, values_only=True), start=2):
        if r[0] is None:
            continue
        year, rank, country, points, progress = r
        year, rank, points = to_int(year), to_int(rank), to_num(points)
        if year is None or rank is None or not country or points is None:
            fr_skipped.append((i, r))
            continue
        name, code = resolve(country)
        final_rank.append({
            "year": year, "rank": rank, "country": name, "country_code": code,
            "points": points,
            "progress": None if progress is None else str(progress).strip(),
        })

    return categories, entries, skipped, final_rank, fr_skipped


def read_final_ranking(resolve):
    """FINALRANKING<year>.xlsx → latest-sports metadata + the A–E snapshot
    (used only to cross-check what the dashboard will compute)."""
    ws = load_workbook(FINAL_FILE)["Final Ranking"]
    latest_sports, sports_included, snapshot = [], None, {}
    for r in ws.iter_rows(min_row=2, max_col=7, values_only=True):
        if r[1] and r[2] is not None:
            name, _ = resolve(r[1])
            snapshot[name] = (to_int(r[0]), to_num(r[2]), to_int(r[4]))
        if r[5]:
            latest_sports.append(" ".join(str(r[5]).split()).upper())
        if r[6] is not None and sports_included is None:
            sports_included = to_int(r[6])
    return latest_sports, sports_included, snapshot


# ── Validation ────────────────────────────────────────────────────────────────
def validate(categories, entries, final_rank, latest_sports, sports_included, snapshot):
    problems = []
    cat_names = {c["sport"] for c in categories}
    missing = sorted({e["sport"] for e in entries} - cat_names)
    if missing:
        problems.append(f"Sports in BY COUNTRY missing from CATEGORIES: {missing}")
    bad_latest = sorted(set(latest_sports) - cat_names)
    if bad_latest:
        problems.append(f"Latest-sports list has unknown sports: {bad_latest}")

    current_year = max(e["year"] for e in entries)
    cur = [e for e in entries if e["year"] == current_year]
    n_sports = len({e["sport"] for e in cur})
    if sports_included is not None and n_sports != sports_included:
        problems.append(f"{current_year}: workbook says {sports_included} sports included, "
                        f"BY COUNTRY has {n_sports}")

    totals = collections.defaultdict(float)
    for e in cur:
        totals[e["country"]] += e["points"]
    drift = [(c, round(snap[1], 2), round(totals.get(c, 0), 2))
             for c, snap in snapshot.items() if abs(snap[1] - totals.get(c, 0)) > 0.01]
    if drift:
        problems.append(f"{len(drift)} countries' {current_year} totals differ between "
                        f"FINALRANKING and BY COUNTRY, e.g. {drift[:5]}")
    return current_year, problems


# ── Upload ────────────────────────────────────────────────────────────────────
def get_client():
    if not SUPABASE_KEY:
        print("ERROR: set SUPABASE_SERVICE_KEY (the service_role secret key).")
        sys.exit(1)
    try:
        from supabase import create_client
    except ImportError:
        print("ERROR: supabase package not installed. Run:  .venv/bin/pip install supabase")
        sys.exit(1)
    return create_client(SUPABASE_URL, SUPABASE_KEY)


def dashboard_edits_since_last_import(client):
    res = client.table("ds_settings").select("value").eq("key", "last_import").execute()
    last = res.data[0]["value"]["at"] if res.data else None
    q = client.table("ds_audit_log").select("id", count="exact").neq("action", "import")
    if last:
        q = q.gt("at", last)
    return q.limit(1).execute().count or 0


def insert_batches(client, table, rows):
    for i in range(0, len(rows), BATCH_SIZE):
        client.table(table).insert(rows[i:i + BATCH_SIZE]).execute()
        print(f"  {table}: {min(i + BATCH_SIZE, len(rows))}/{len(rows)}", end="\r")
    print(f"  ✓ {table} ({len(rows)} rows)              ")


def upload(client, countries, categories, entries, final_rank, latest_sports, current_year):
    # Children first, then parents (foreign keys). PostgREST needs a filter on
    # delete, so use one that matches every row.
    print("\nClearing ds_* tables…")
    client.table("ds_entries").delete().gte("id", 0).execute()
    client.table("ds_final_rank").delete().gte("id", 0).execute()
    client.table("ds_categories").delete().neq("sport", "").execute()
    client.table("ds_countries").delete().neq("code", "").execute()

    print("Uploading…")
    insert_batches(client, "ds_countries", countries)
    insert_batches(client, "ds_categories", categories)
    insert_batches(client, "ds_final_rank", final_rank)
    insert_batches(client, "ds_entries", entries)

    now = dt.datetime.now(dt.timezone.utc).isoformat()
    client.table("ds_settings").upsert([
        {"key": "latest_sports", "value": {"year": current_year, "sports": latest_sports}, "updated_at": now},
        {"key": "last_import", "value": {"at": now, "source": STORAGE_FILE.name}, "updated_at": now},
    ]).execute()
    client.table("ds_audit_log").insert({
        "username": "import", "action": "import", "table_name": "ds_*",
        "after": {"entries": len(entries), "final_rank": len(final_rank), "source": STORAGE_FILE.name},
    }).execute()


# ── Main ──────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    dry_run = "--dry-run" in sys.argv
    force = "--force" in sys.argv

    print("WSR Admin Data Import")
    print("=====================")
    print(f"Data directory: {DATA_DIR}")

    countries = read_countries()
    resolve = CountryResolver(countries)
    categories, entries, skipped, final_rank, fr_skipped = read_storage(resolve)
    latest_sports, sports_included, snapshot = read_final_ranking(resolve)
    current_year, problems = validate(categories, entries, final_rank, latest_sports,
                                      sports_included, snapshot)

    years = collections.Counter(e["year"] for e in entries)
    print(f"\nCountries:   {len(countries)}")
    print(f"Categories:  {len(categories)}  {dict(collections.Counter(c['type'] for c in categories))}")
    print(f"Entries:     {len(entries)}  ({min(years)}–{max(years)})")
    print(f"Final rank:  {len(final_rank)}  ({min(r['year'] for r in final_rank)}–{max(r['year'] for r in final_rank)})")
    print(f"Current year {current_year}: {years[current_year]} entries; latest sports added: {latest_sports}")
    if skipped or fr_skipped:
        print(f"⚠ Skipped {len(skipped)} BY COUNTRY / {len(fr_skipped)} FINAL RANK rows with missing values:")
        for row in (skipped + fr_skipped)[:10]:
            print("   ", row)
    if resolve.unmatched:
        print(f"⚠ Names not in COUNTRIESLIST (kept, no country code): {dict(resolve.unmatched)}")
    for p in problems:
        print(f"⚠ {p}")

    if dry_run:
        print("\nDry run — nothing uploaded.")
        sys.exit(0)

    client = get_client()
    edits = dashboard_edits_since_last_import(client)
    if edits and not force:
        print(f"\nSTOP: {edits} change(s) were made in the dashboard since the last import.")
        print("Re-importing would overwrite them. Export them first, then re-run with --force.")
        sys.exit(1)

    upload(client, countries, categories, entries, final_rank, latest_sports, current_year)
    print("\n✅ Admin data imported.")
